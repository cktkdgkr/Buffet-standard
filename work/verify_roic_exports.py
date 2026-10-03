"""
Evaluate every formula in the ROIC workbook and check it against roic.json.

Recalculating only proves the formulas evaluate without error. A formula
pointing one row off evaluates perfectly and returns a number for the wrong
company. So each evaluated result is compared to the figure analyse_roic.py
produced, which is what the workbook claims to show.

Exit code is non-zero if anything disagrees.

    python3 verify_roic_exports.py
"""

import json
import os
import re
import sys

from openpyxl import load_workbook

WORK = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, WORK)
from verify_exports import Evaluator, close                    # noqa: E402

XLSX = os.path.join(WORK, "미국50개사_ROIC_WACC비교.xlsx")
ROIC = os.path.join(WORK, "roic", "roic.json")
YEARS = list(range(2016, 2026))


def main():
    data = json.load(open(ROIC))
    wb = load_workbook(XLSX, data_only=False)
    ev = Evaluator(wb)

    by = {c["ticker"]: {y["fiscal_year"]: y for y in c["years"]}
          for c in data["companies"]}
    summ = {c["ticker"]: c["summary"] for c in data["companies"]}

    checks = 0
    fails = []

    def expect(sheet, coord, want, label, rel=2e-6):
        nonlocal checks
        checks += 1
        try:
            got = ev.value(sheet, coord)
        except Exception as exc:                               # noqa: BLE001
            fails.append(f"{label} [{sheet}!{coord}]: 평가 실패 {exc}")
            return
        if want is None:
            if got not in (None, "", "—"):
                fails.append(f"{label} [{sheet}!{coord}]: 값이 있어야 "
                             f"하지 않음 — {got!r}")
            return
        if not isinstance(got, (int, float)):
            fails.append(f"{label} [{sheet}!{coord}]: 숫자가 아님 {got!r}")
            return
        if not close(got, want, rel):
            fails.append(f"{label} [{sheet}!{coord}]: {got!r} != {want!r}")

    # --- 계산 명세: every derived column, every row -----------------------
    ws = wb["계산 명세"]
    rowmap = {}
    for r in range(2, ws.max_row + 1):
        tk = ws.cell(row=r, column=2).value
        fy = ws.cell(row=r, column=3).value
        if not tk or not fy:
            continue
        rowmap[(tk, fy)] = r
        y = by[tk][fy]
        tag = f"{tk} FY{fy}"

        expect("계산 명세", f"J{r}", y["nopat"], f"{tag} NOPAT")
        expect("계산 명세", f"N{r}", y["invested_capital"],
               f"{tag} 투하자본")
        if fy == 2015:
            # Reference row: it supplies FY2016's opening capital. Its own
            # average would need FY2014, so the ratio cells must be blank.
            for col in "OPQRSTUV":
                checks += 1
                if ws.cell(row=r, column=ord(col) - 64).value != "—":
                    fails.append(f"{tag} {col}{r}: 기초 참조 행인데 값이 있음")
            continue
        expect("계산 명세", f"P{r}", y["invested_capital_avg"],
               f"{tag} 평균 투하자본")
        expect("계산 명세", f"R{r}", y["roic"], f"{tag} ROIC")
        if y["roic"] is not None:
            expect("계산 명세", f"S{r}", y["roic"] - data["wacc"],
                   f"{tag} 스프레드")
        expect("계산 명세", f"T{r}", y["capital_intensity"],
               f"{tag} 분모/매출")

        # The prior-capital cell must point at the prior year's row, not merely
        # hold the right number: a company whose first window year is FY2016
        # has no prior row and must show a dash.
        prev = rowmap.get((tk, fy - 1))
        raw = str(ws.cell(row=r, column=15).value)
        checks += 1
        if prev:
            if raw != f"=N{prev}":
                fails.append(f"{tag} 직전 투하자본 참조: {raw} != =N{prev}")
        elif raw != "—":
            fails.append(f"{tag} 직전 투하자본: 직전 행이 없는데 {raw}")

    # Every window year of every company must appear exactly once
    want_rows = {(c["ticker"], y["fiscal_year"])
                 for c in data["companies"] for y in c["years"]
                 if 2015 <= y["fiscal_year"] <= 2025}
    checks += 1
    if set(rowmap) != want_rows:
        missing = want_rows - set(rowmap)
        extra = set(rowmap) - want_rows
        fails.append(f"계산 명세 행 집합 불일치 — 누락 {sorted(missing)[:5]}, "
                     f"초과 {sorted(extra)[:5]}")

    # --- the two year grids ---------------------------------------------
    for sheet, offset in (("ROIC 연도별", 0.0), ("ROIC−WACC 스프레드",
                                              -data["wacc"])):
        ws = wb[sheet]
        seen = set()
        for r in range(3, ws.max_row + 1):
            tk = ws.cell(row=r, column=3).value
            if not tk or tk not in by:
                continue
            seen.add((sheet, tk))
            for i, fy in enumerate(YEARS):
                coord = f"{chr(ord('D') + i)}{r}"
                y = by[tk].get(fy)
                want = None if (y is None or y["roic"] is None) \
                    else y["roic"] + offset
                expect(sheet, coord, want, f"{sheet} {tk} FY{fy}")

            roics = [y["roic"] + offset for fy in YEARS
                     if (y := by[tk].get(fy)) and y["roic"] is not None]
            med_col = chr(ord('D') + len(YEARS))
            if roics:
                s = summ[tk]
                expect(sheet, f"{med_col}{r}", s["roic_median"] + offset,
                       f"{sheet} {tk} 중위")
                expect(sheet, f"{chr(ord(med_col) + 1)}{r}",
                       s["roic_min"] + offset, f"{sheet} {tk} 최소")
                expect(sheet, f"{chr(ord(med_col) + 2)}{r}",
                       s["roic_max"] + offset, f"{sheet} {tk} 최대")
                expect(sheet, f"{chr(ord(med_col) + 3)}{r}",
                       float(s["years_above_wacc"]),
                       f"{sheet} {tk} 초과 연수")
                expect(sheet, f"{chr(ord(med_col) + 4)}{r}",
                       float(s["years_usable"]), f"{sheet} {tk} 산출 연수")

        checks += 1
        want_tk = {(sheet, c["ticker"]) for c in data["companies"]}
        if seen != want_tk:
            fails.append(f"{sheet}: 기업 집합 불일치 — "
                         f"누락 {sorted(t for _, t in want_tk - seen)[:6]}")

    # --- sheets whose derived cells are simple differences ---------------
    # The expected value is read out of each cell's own formula rather than
    # assumed, because one sheet carries two blocks with opposite sign
    # conventions: the ROIC correction block subtracts before from after, and
    # the net-cash block subtracts the recomputed figure from the one the study
    # used, so the error reads positive when the study overstated net cash.
    diff = re.compile(r"^=([A-Z]+)(\d+)-([A-Z]+)(\d+)$")
    for sheet in ("구 조사 대조", "정정 내역", "경로 검정"):
        ws = wb[sheet]
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                m = diff.match(cell.value.replace(" ", ""))
                if not m:
                    continue
                a, ar, b, br = m.groups()
                va = ev.value(sheet, f"{a}{ar}")
                vb = ev.value(sheet, f"{b}{br}")
                if isinstance(va, (int, float)) and isinstance(vb, (int, float)):
                    expect(sheet, cell.coordinate, va - vb,
                           f"{sheet} {cell.coordinate} 차이")

    print(f"수식 검사 {checks}건 중 통과 {checks - len(fails)}건, "
          f"실패 {len(fails)}건")
    for f in fails[:30]:
        print(f"  실패: {f}")
    if len(fails) > 30:
        print(f"  … 외 {len(fails) - 30}건")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
